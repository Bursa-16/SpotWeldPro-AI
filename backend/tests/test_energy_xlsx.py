"""
test_energy_xlsx.py — XLSX import tests for P1-HOTFIX-01.

Test IDs:
  XLSX-01   valid workbook, all columns
  XLSX-02   missing required column (point_code absent)
  XLSX-03   invalid numeric value
  XLSX-04   duplicate point codes within the workbook
  XLSX-05   formula injection string in a cell
  XLSX-06   workbook with multiple sheets — uses active (first)
  XLSX-07   empty workbook (no data rows, header only)
  XLSX-08   file too large (>10 MB)
  XLSX-09   non-XLSX bytes (truncated / corrupt)
  XLSX-10   existing-point-code duplicate detection
  CSV-COMPAT CSV backward compat still works through validate_csv_import

API-XLSX-01  POST /import/preview/file — valid XLSX upload
API-XLSX-02  POST /import/preview/file — XLSX injection blocked
API-XLSX-03  POST /import/preview/file — valid CSV upload
API-XLSX-04  POST /import/preview/file — rejects without auth
"""

from __future__ import annotations

import io
import pytest
import openpyxl

from app.domain.excel_import import validate_xlsx_import, MAX_XLSX_BYTES
from app.domain.import_validation import validate_csv_import


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

HEADERS = ["point_code", "group_name", "e_current_kj", "e_min_kj", "e_ideal_kj",
           "e_max_kj", "validation_note"]


def _make_xlsx(rows: list[list], sheet_name: str = "Sheet1") -> bytes:
    """Build an in-memory XLSX workbook with one sheet."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _make_multi_sheet_xlsx(sheets: dict[str, list[list]]) -> bytes:
    """Build an XLSX with multiple named sheets."""
    wb = openpyxl.Workbook()
    first = True
    for name, rows in sheets.items():
        if first:
            ws = wb.active
            ws.title = name
            first = False
        else:
            ws = wb.create_sheet(title=name)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# XLSX-01: valid workbook with two accepted rows
# ---------------------------------------------------------------------------
def test_xlsx01_valid_workbook():
    data = [
        HEADERS,
        ["WP-001", "Group-A", 0.400, 0.280, 0.320, 0.420, "Validated"],
        ["WP-002", "Group-A", 0.380, 0.280, 0.320, 0.420, ""],
    ]
    result = validate_xlsx_import(_make_xlsx(data))
    assert result.total_rows == 2
    assert result.accepted_count == 2
    assert result.rejected_count == 0
    assert not result.has_blocking_errors
    assert not result.formula_injection_detected


# ---------------------------------------------------------------------------
# XLSX-02: missing required column (no point_code column)
# ---------------------------------------------------------------------------
def test_xlsx02_missing_required_column():
    data = [
        ["group_name", "e_current_kj"],
        ["Group-A", 0.400],
    ]
    result = validate_xlsx_import(_make_xlsx(data))
    assert "point_code" in result.missing_columns
    assert result.has_blocking_errors is True


# ---------------------------------------------------------------------------
# XLSX-03: invalid numeric value
# ---------------------------------------------------------------------------
def test_xlsx03_invalid_numeric():
    data = [
        HEADERS,
        ["WP-001", "Group-A", "NOTANUMBER", 0.280, 0.320, 0.420, ""],
    ]
    result = validate_xlsx_import(_make_xlsx(data))
    assert result.rejected_count == 1
    assert any(e.error_type == "INVALID_NUMBER" for e in result.errors)


# ---------------------------------------------------------------------------
# XLSX-04: duplicate point codes within the workbook
# ---------------------------------------------------------------------------
def test_xlsx04_duplicate_within_file():
    data = [
        HEADERS,
        ["WP-001", "Group-A", 0.400, 0.280, 0.320, 0.420, ""],
        ["WP-001", "Group-B", 0.380, 0.280, 0.320, 0.420, ""],
    ]
    result = validate_xlsx_import(_make_xlsx(data))
    assert "WP-001" in result.duplicate_point_codes
    assert result.has_blocking_errors is True


# ---------------------------------------------------------------------------
# XLSX-05: formula injection — text cell starting with injection character
#
# Note on Excel formula cells: openpyxl writes "=CMD()" as an actual formula
# cell. Reading it back with data_only=True returns None (no cached result),
# which becomes an empty string → MISSING_REQUIRED rather than FORMULA_INJECTION.
# That is the correct security outcome — formulas are never evaluated.
# This test uses "|" prefix, which openpyxl stores as a text string, verifying
# that injection chars in text cells are still caught.
# ---------------------------------------------------------------------------
def test_xlsx05_formula_injection():
    data = [
        HEADERS,
        ["|pipe-injection", "Group-A", 0.400, 0.280, 0.320, 0.420, ""],
    ]
    result = validate_xlsx_import(_make_xlsx(data))
    assert result.formula_injection_detected is True
    assert result.has_blocking_errors is True
    assert result.accepted_count == 0


def test_xlsx05c_excel_formula_cells_become_missing():
    """Excel formula cells (=...) read as None with data_only=True → MISSING_REQUIRED.

    This is the correct security behavior: formulas are never evaluated.
    """
    data = [
        HEADERS,
        ["=CMD()", "Group-A", 0.400, 0.280, 0.320, 0.420, ""],
    ]
    result = validate_xlsx_import(_make_xlsx(data))
    assert any(e.error_type == "MISSING_REQUIRED" for e in result.errors)
    assert result.accepted_count == 0
    assert result.formula_injection_detected is False


# ---------------------------------------------------------------------------
# XLSX-05b: DDE injection pattern
# ---------------------------------------------------------------------------
def test_xlsx05b_dde_injection():
    data = [
        HEADERS,
        ["DDE(cmd,/C,calc)", "Group-A", 0.400, 0.280, 0.320, 0.420, ""],
    ]
    result = validate_xlsx_import(_make_xlsx(data))
    assert result.formula_injection_detected is True
    assert result.accepted_count == 0


# ---------------------------------------------------------------------------
# XLSX-06: multiple sheets — uses active (first) sheet
# ---------------------------------------------------------------------------
def test_xlsx06_multiple_sheets_uses_first():
    sheets = {
        "Energy": [
            HEADERS,
            ["WP-001", "Group-A", 0.400, 0.280, 0.320, 0.420, ""],
        ],
        "Ignored": [
            HEADERS,
            ["WP-BAD", "Group-X", "NOTANUMBER", None, None, None, ""],
        ],
    }
    result = validate_xlsx_import(_make_multi_sheet_xlsx(sheets))
    assert result.accepted_count == 1
    assert result.total_rows == 1
    assert not result.formula_injection_detected


# ---------------------------------------------------------------------------
# XLSX-06b: select a specific sheet by name
# ---------------------------------------------------------------------------
def test_xlsx06b_select_sheet_by_name():
    sheets = {
        "Summary": [
            HEADERS,
            ["WP-S", "Grp", 0.4, 0.28, 0.32, 0.42, ""],
        ],
        "Detail": [
            HEADERS,
            ["WP-D1", "Grp", 0.38, 0.28, 0.32, 0.42, ""],
            ["WP-D2", "Grp", 0.41, 0.28, 0.32, 0.42, ""],
        ],
    }
    xlsx = _make_multi_sheet_xlsx(sheets)
    result = validate_xlsx_import(xlsx, sheet_name="Detail")
    assert result.total_rows == 2
    assert result.accepted_count == 2


# ---------------------------------------------------------------------------
# XLSX-07: header-only workbook (no data rows)
# ---------------------------------------------------------------------------
def test_xlsx07_empty_data_rows():
    data = [HEADERS]
    result = validate_xlsx_import(_make_xlsx(data))
    assert result.total_rows == 0
    # No blocking errors from an empty-but-valid file
    assert result.accepted_count == 0


# ---------------------------------------------------------------------------
# XLSX-08: file too large
# ---------------------------------------------------------------------------
def test_xlsx08_file_too_large():
    oversized = b"x" * (MAX_XLSX_BYTES + 1)
    result = validate_xlsx_import(oversized)
    assert result.has_blocking_errors is True
    assert any(e.error_type == "FILE_TOO_LARGE" for e in result.errors)


# ---------------------------------------------------------------------------
# XLSX-09: corrupt / non-XLSX bytes
# ---------------------------------------------------------------------------
def test_xlsx09_corrupt_bytes():
    result = validate_xlsx_import(b"This is not an XLSX file at all.")
    assert result.has_blocking_errors is True
    assert any(e.error_type == "PARSE_ERROR" for e in result.errors)


# ---------------------------------------------------------------------------
# XLSX-10: existing-point-code duplicate detection
# ---------------------------------------------------------------------------
def test_xlsx10_existing_code_duplicate():
    data = [
        HEADERS,
        ["WP-DB-01", "Group-A", 0.400, 0.280, 0.320, 0.420, ""],
    ]
    result = validate_xlsx_import(
        _make_xlsx(data),
        existing_point_codes={"WP-DB-01"},
    )
    assert "WP-DB-01" in result.duplicate_point_codes
    assert result.has_blocking_errors is True


# ---------------------------------------------------------------------------
# CSV-COMPAT: CSV backward compatibility
# ---------------------------------------------------------------------------
def test_csv_compat_still_works():
    csv = (
        "point_code,group_name,e_current_kj,e_min_kj,e_ideal_kj,e_max_kj,validation_note\n"
        "WP-001,Group-A,0.400,0.280,0.320,0.420,OK\n"
    )
    result = validate_csv_import(csv)
    assert result.accepted_count == 1
    assert not result.has_blocking_errors


# ---------------------------------------------------------------------------
# XLSX window ordering: e_min > e_max rejected
# ---------------------------------------------------------------------------
def test_xlsx_window_ordering_rejected():
    data = [
        HEADERS,
        ["WP-001", "Group-A", 0.400, 0.500, 0.450, 0.300, ""],
    ]
    result = validate_xlsx_import(_make_xlsx(data))
    assert any(e.error_type == "INVALID_NUMBER" for e in result.errors)


# ---------------------------------------------------------------------------
# XLSX numeric cells: native float cells round-trip correctly
# ---------------------------------------------------------------------------
def test_xlsx_native_float_cells():
    """Excel cells with float type (not string) parse correctly."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(HEADERS)
    ws["A2"] = "WP-FLOAT"
    ws["B2"] = "Group-F"
    ws["C2"] = 0.35   # native float cell
    ws["D2"] = 0.28
    ws["E2"] = 0.32
    ws["F2"] = 0.42
    ws["G2"] = ""
    buf = io.BytesIO()
    wb.save(buf)
    result = validate_xlsx_import(buf.getvalue())
    assert result.accepted_count == 1
    row = result.preview_rows[0]
    assert row.e_current_kj == pytest.approx(0.35)
    assert row.e_ideal_kj == pytest.approx(0.32)


# ---------------------------------------------------------------------------
# API tests: /import/preview/file
# ---------------------------------------------------------------------------

def test_api_xlsx01_valid_xlsx_upload(client, auth_headers):
    data = [
        HEADERS,
        ["WP-API", "Group-A", 0.400, 0.280, 0.320, 0.420, "API test"],
    ]
    xlsx_bytes = _make_xlsx(data)
    resp = client.post(
        "/api/v1/energy/import/preview/file",
        files={"file": ("test.xlsx", xlsx_bytes,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=auth_headers,
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["accepted_count"] == 1
    assert not body["formula_injection_detected"]
    assert not body["has_blocking_errors"]


def test_api_xlsx02_xlsx_injection_blocked(client, auth_headers):
    data = [
        HEADERS,
        ["|EVIL|", "Group-A", 0.400, 0.280, 0.320, 0.420, ""],
    ]
    xlsx_bytes = _make_xlsx(data)
    resp = client.post(
        "/api/v1/energy/import/preview/file",
        files={"file": ("test.xlsx", xlsx_bytes,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        headers=auth_headers,
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["formula_injection_detected"] is True
    assert body["has_blocking_errors"] is True
    assert body["accepted_count"] == 0


def test_api_xlsx03_csv_upload_via_file_endpoint(client, auth_headers):
    csv_text = (
        "point_code,group_name,e_current_kj,e_min_kj,e_ideal_kj,e_max_kj,validation_note\n"
        "WP-CSV,Group-A,0.400,0.280,0.320,0.420,CSV via file endpoint\n"
    )
    resp = client.post(
        "/api/v1/energy/import/preview/file",
        files={"file": ("test.csv", csv_text.encode(), "text/csv")},
        headers=auth_headers,
    )
    assert resp.status_code == 200
    body = resp.json()
    assert body["accepted_count"] == 1
    assert not body["has_blocking_errors"]


def test_api_xlsx04_requires_auth(client):
    """File upload endpoint requires authentication."""
    xlsx_bytes = _make_xlsx([HEADERS, ["WP-001", "G", 0.4, 0.28, 0.32, 0.42, ""]])
    resp = client.post(
        "/api/v1/energy/import/preview/file",
        files={"file": ("test.xlsx", xlsx_bytes,
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 401
