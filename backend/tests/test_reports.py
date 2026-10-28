"""
test_reports.py — P2-F Engineering Report tests.

Test IDs:
  RPT-01 … RPT-05  CSV generation and content
  RPT-06 … RPT-10  XLSX generation and content
  RPT-11 … RPT-15  PDF generation and content
  RPT-16 … RPT-20  Report calculations and data accuracy
  RPT-21 … RPT-23  NOT_ESTABLISHED handling
  RPT-24 … RPT-27  Security: formula injection prevention
  RPT-28 … RPT-30  Traceability fields
  RPT-31 … RPT-33  Report language and metadata
  RPT-34 … RPT-36  Authorization (API endpoints)
  RPT-37 … RPT-38  Backend regression guard

Run: pytest tests/test_reports.py -v
"""

from __future__ import annotations

import csv
import hashlib
import io
import json

import openpyxl
import pytest

from app.domain.energy_aggregation import WeldPointEnergySpec
from app.domain.report_engine import (
    _safe_cell,
    build_report_context,
    generate_csv_report,
    generate_pdf_report,
    generate_xlsx_report,
)
from app.schemas.reports import EnergyReportRequest


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------

def _specs(points: list[tuple]) -> list[WeldPointEnergySpec]:
    return [
        WeldPointEnergySpec(
            point_code=p[0], group_name=p[1],
            e_current_kj=p[2], e_min_kj=p[3], e_ideal_kj=p[4], e_max_kj=p[5],
        )
        for p in points
    ]


def _typical_specs() -> list[WeldPointEnergySpec]:
    return _specs([
        ("WP-001", "G-A", 0.55, 0.28, 0.32, 0.42),
        ("WP-002", "G-A", 0.40, 0.28, 0.32, 0.42),
        ("WP-003", "G-B", 0.30, 0.28, 0.32, 0.42),
        ("WP-004", "G-B", 0.48, 0.28, 0.32, 0.42),
    ])


def _ctx(**kwargs):
    defaults = dict(production_volume_annual=200_000, tariff_eur_per_kwh=0.18, tariff_usd_per_kwh=0.22)
    defaults.update(kwargs)
    return build_report_context(_typical_specs(), **defaults)


def _api_req() -> dict:
    return {
        "points": [
            {"point_code": "WP-001", "group_name": "G-A",
             "e_current_kj": 0.55, "e_min_kj": 0.28, "e_ideal_kj": 0.32, "e_max_kj": 0.42},
            {"point_code": "WP-002", "group_name": "G-A",
             "e_current_kj": 0.40, "e_min_kj": 0.28, "e_ideal_kj": 0.32, "e_max_kj": 0.42},
        ],
        "production_volume_annual": 100_000,
        "tariff_eur_per_kwh": 0.18,
        "tariff_usd_per_kwh": 0.22,
        "project_id": "PROJ-TEST",
        "passport_id": "DWP-XYZ",
        "report_language": "en",
    }


# ---------------------------------------------------------------------------
# RPT-01 … RPT-05: CSV
# ---------------------------------------------------------------------------

def test_rpt01_csv_is_valid_utf8_with_bom():
    """RPT-01: CSV output starts with UTF-8 BOM."""
    ctx = _ctx()
    data = generate_csv_report(ctx)
    assert data[:3] == b"\xef\xbb\xbf"


def test_rpt02_csv_contains_report_id():
    """RPT-02: CSV contains the report ID."""
    ctx = _ctx()
    data = generate_csv_report(ctx)
    text = data.decode("utf-8-sig")
    assert ctx.report_id in text


def test_rpt03_csv_weld_point_codes_present():
    """RPT-03: All point codes appear in CSV output."""
    ctx = _ctx()
    data = generate_csv_report(ctx)
    text = data.decode("utf-8-sig")
    for row in ctx.aggr.rows:
        assert row.spec.point_code in text


def test_rpt04_csv_has_disclaimer():
    """RPT-04: Disclaimer text is present in CSV."""
    ctx = _ctx()
    data = generate_csv_report(ctx)
    text = data.decode("utf-8-sig")
    assert "THEORETICAL" in text
    assert "NOT a verified electricity" in text


def test_rpt05_csv_financial_section_when_tariff_given():
    """RPT-05: Financial section appears when tariffs are provided."""
    ctx = _ctx(tariff_eur_per_kwh=0.20, tariff_usd_per_kwh=0.25, production_volume_annual=500_000)
    data = generate_csv_report(ctx)
    text = data.decode("utf-8-sig")
    assert "Annual Cost EUR" in text


# ---------------------------------------------------------------------------
# RPT-06 … RPT-10: XLSX
# ---------------------------------------------------------------------------

def test_rpt06_xlsx_is_valid_zip():
    """RPT-06: XLSX bytes begin with PK (valid ZIP/OOXML)."""
    ctx = _ctx()
    data = generate_xlsx_report(ctx)
    assert data[:4] == b"PK\x03\x04"


def test_rpt07_xlsx_has_required_sheets():
    """RPT-07: XLSX contains expected sheet names."""
    ctx = _ctx()
    wb = openpyxl.load_workbook(io.BytesIO(generate_xlsx_report(ctx)))
    names = wb.sheetnames
    for expected in ["Summary", "Weld Points", "Groups", "Pareto", "Financial", "Traceability"]:
        assert expected in names, f"Sheet '{expected}' missing"


def test_rpt08_xlsx_weld_points_sheet_row_count():
    """RPT-08: Weld Points sheet has one data row per point plus header."""
    ctx = _ctx()
    wb = openpyxl.load_workbook(io.BytesIO(generate_xlsx_report(ctx)))
    ws = wb["Weld Points"]
    n_data = sum(1 for r in ws.iter_rows(min_row=2) if r[0].value is not None)
    assert n_data == len(ctx.aggr.rows)


def test_rpt09_xlsx_report_id_in_traceability():
    """RPT-09: Report ID appears in Traceability sheet."""
    ctx = _ctx()
    wb = openpyxl.load_workbook(io.BytesIO(generate_xlsx_report(ctx)))
    ws = wb["Traceability"]
    values = [str(c.value or "") for row in ws.iter_rows() for c in row]
    assert ctx.report_id in values


def test_rpt10_xlsx_financial_sheet_has_sensitivity():
    """RPT-10: Financial sheet contains sensitivity matrix rows."""
    ctx = _ctx(production_volume_annual=200_000, tariff_eur_per_kwh=0.18)
    wb = openpyxl.load_workbook(io.BytesIO(generate_xlsx_report(ctx)))
    ws = wb["Financial"]
    # Sensitivity section header or first tariff label must appear
    values = [str(c.value or "") for row in ws.iter_rows() for c in row]
    has_sensitivity = any("EUR/kWh" in v or "Sensitivity" in v for v in values)
    assert has_sensitivity


# ---------------------------------------------------------------------------
# RPT-11 … RPT-15: PDF
# ---------------------------------------------------------------------------

def test_rpt11_pdf_starts_with_header():
    """RPT-11: PDF bytes start with %PDF signature."""
    ctx = _ctx()
    data = generate_pdf_report(ctx)
    assert data[:4] == b"%PDF"


def test_rpt12_pdf_non_empty():
    """RPT-12: PDF is larger than 1 kB."""
    ctx = _ctx()
    data = generate_pdf_report(ctx)
    assert len(data) > 1024


def test_rpt13_pdf_not_established_no_crash():
    """RPT-13: PDF generation succeeds with all NOT_ESTABLISHED points."""
    ne_specs = _specs([
        ("WP-A", "G", None, None, None, None),
        ("WP-B", "G", None, None, None, None),
    ])
    ctx = build_report_context(ne_specs)
    data = generate_pdf_report(ctx)
    assert data[:4] == b"%PDF"


def test_rpt14_pdf_turkish_metadata_no_crash():
    """RPT-14: Turkish project name and analyst name do not crash PDF generation."""
    ctx = build_report_context(
        _typical_specs(),
        project_name="Enerji Analiz Raporu — Şase Grubu",
        analyst_name="Mühendis Çelik",
    )
    data = generate_pdf_report(ctx)
    assert data[:4] == b"%PDF"


def test_rpt15_pdf_valid_opaque_bytes():
    """RPT-15: PDF bytes end with %%EOF marker."""
    ctx = _ctx()
    data = generate_pdf_report(ctx)
    assert b"%%EOF" in data[-256:]


# ---------------------------------------------------------------------------
# RPT-16 … RPT-20: Calculation accuracy
# ---------------------------------------------------------------------------

def test_rpt16_excess_above_max_point():
    """RPT-16: ABOVE_MAX point (0.55 vs 0.42 max) has positive excess."""
    ctx = _ctx()
    above = [r for r in ctx.aggr.rows if r.spec.point_code == "WP-001"]
    assert len(above) == 1
    assert above[0].window.e_excess_potential_kj is not None
    assert above[0].window.e_excess_potential_kj > 0


def test_rpt17_in_window_point_zero_excess():
    """RPT-17: IN_WINDOW point (0.40 kJ in [0.28..0.42]) — excess ≥ 0."""
    ctx = _ctx()
    inwin = [r for r in ctx.aggr.rows if r.spec.point_code == "WP-002"]
    assert len(inwin) == 1
    # WP-002 (0.40) > E_ideal (0.32) → positive excess; no negative excess
    exc = inwin[0].window.e_excess_potential_kj
    assert exc is not None
    assert exc >= 0


def test_rpt18_financial_kj_to_kwh_conversion():
    """RPT-18: kJ → kWh: part_excess_kwh = part_excess_kj / 3600."""
    ctx = _ctx()
    f = ctx.extras.financial
    if f.part_excess_kj is not None and f.part_excess_kwh is not None:
        assert abs(f.part_excess_kwh - f.part_excess_kj / 3600.0) < 1e-10


def test_rpt19_csv_pareto_rank1_has_highest_excess():
    """RPT-19: First Pareto row has the highest excess value."""
    ctx = _ctx()
    pareto = ctx.extras.pareto_detailed
    sorted_excess = sorted(
        [r.window.e_excess_potential_kj for r in ctx.aggr.rows
         if r.window.e_excess_potential_kj is not None],
        reverse=True,
    )
    if sorted_excess and pareto[0].e_excess_kj is not None:
        assert abs(pareto[0].e_excess_kj - sorted_excess[0]) < 1e-10


def test_rpt20_annual_cost_scales_linearly_with_volume():
    """RPT-20: Doubling annual production volume doubles annual cost EUR."""
    ctx1 = build_report_context(
        _typical_specs(),
        production_volume_annual=100_000,
        tariff_eur_per_kwh=0.20,
        tariff_usd_per_kwh=0.24,
    )
    ctx2 = build_report_context(
        _typical_specs(),
        production_volume_annual=200_000,
        tariff_eur_per_kwh=0.20,
        tariff_usd_per_kwh=0.24,
    )
    f1 = ctx1.extras.financial
    f2 = ctx2.extras.financial
    if f1.annual_cost_eur and f2.annual_cost_eur:
        assert abs(f2.annual_cost_eur / f1.annual_cost_eur - 2.0) < 1e-6


# ---------------------------------------------------------------------------
# RPT-21 … RPT-23: NOT_ESTABLISHED handling
# ---------------------------------------------------------------------------

def test_rpt21_not_established_status_preserved_in_csv():
    """RPT-21: NOT_ESTABLISHED points appear with correct status in CSV."""
    ne_specs = _specs([("WP-NE", "G", None, None, None, None)])
    ctx = build_report_context(ne_specs)
    data = generate_csv_report(ctx)
    text = data.decode("utf-8-sig")
    assert "NOT_ESTABLISHED" in text


def test_rpt22_not_established_financial_is_none():
    """RPT-22: All NOT_ESTABLISHED → financial kpis are None/absent."""
    ne_specs = _specs([
        ("A", "G", None, None, None, None),
        ("B", "G", None, None, None, None),
    ])
    ctx = build_report_context(ne_specs, tariff_eur_per_kwh=0.18, production_volume_annual=100_000)
    f = ctx.extras.financial
    assert f.annual_cost_eur is None
    assert f.part_excess_kj is None


def test_rpt23_not_established_xlsx_no_fabricated_values():
    """RPT-23: XLSX rows for NOT_ESTABLISHED points contain '--' not computed values."""
    ne_specs = _specs([("WP-NE", "G", None, None, None, None)])
    ctx = build_report_context(ne_specs)
    wb = openpyxl.load_workbook(io.BytesIO(generate_xlsx_report(ctx)))
    ws = wb["Weld Points"]
    row2 = [c.value for c in ws[2]]  # first data row
    # e_min, e_ideal, e_max should be None or '--', not computed numbers
    for v in row2[3:6]:  # columns D-F: E_min, E_ideal, E_max
        assert v is None or v == "--"


# ---------------------------------------------------------------------------
# RPT-24 … RPT-27: Security — formula injection prevention
# ---------------------------------------------------------------------------

def test_rpt24_safe_cell_escapes_equals():
    """RPT-24: =CMD() injection is prefixed with apostrophe."""
    result = _safe_cell("=CMD(malicious)")
    assert result.startswith("'")
    assert "CMD" in result


def test_rpt25_safe_cell_escapes_plus():
    """RPT-25: + prefix injection is escaped."""
    assert _safe_cell("+malicious").startswith("'")


def test_rpt26_safe_cell_escapes_at():
    """RPT-26: @ prefix injection is escaped."""
    assert _safe_cell("@malicious").startswith("'")


def test_rpt27_csv_injection_payload_in_point_code_is_escaped():
    """RPT-27: Formula injection in point_code is escaped in CSV output."""
    injection_specs = _specs([("=SUM(1,2)", "G", 0.40, 0.28, 0.32, 0.42)])
    ctx = build_report_context(injection_specs)
    data = generate_csv_report(ctx)
    text = data.decode("utf-8-sig")
    # The raw formula should not appear unescaped; it should be prefixed
    assert ",'=SUM(1,2)'" in text or "\"'=SUM(1,2)\"" in text or "'=SUM(1,2)" in text


# ---------------------------------------------------------------------------
# RPT-28 … RPT-30: Traceability
# ---------------------------------------------------------------------------

def test_rpt28_report_id_format():
    """RPT-28: Report ID begins with SWP-RPT- prefix."""
    ctx = _ctx()
    assert ctx.report_id.startswith("SWP-RPT-")


def test_rpt29_input_hash_deterministic():
    """RPT-29: Same inputs produce the same input hash."""
    req1 = EnergyReportRequest(**{
        "points": [{"point_code": "A", "group_name": "G", "e_current_kj": 0.5,
                    "e_min_kj": 0.2, "e_ideal_kj": 0.35, "e_max_kj": 0.6}],
        "tariff_eur_per_kwh": 0.18,
    })
    req2 = EnergyReportRequest(**{
        "points": [{"point_code": "A", "group_name": "G", "e_current_kj": 0.5,
                    "e_min_kj": 0.2, "e_ideal_kj": 0.35, "e_max_kj": 0.6}],
        "tariff_eur_per_kwh": 0.18,
    })
    assert req1.input_hash == req2.input_hash


def test_rpt30_different_inputs_produce_different_hash():
    """RPT-30: Different inputs produce different hashes."""
    req1 = EnergyReportRequest(points=[
        {"point_code": "A", "group_name": "G", "e_current_kj": 0.5,
         "e_min_kj": 0.2, "e_ideal_kj": 0.35, "e_max_kj": 0.6}
    ])
    req2 = EnergyReportRequest(points=[
        {"point_code": "A", "group_name": "G", "e_current_kj": 0.6,
         "e_min_kj": 0.2, "e_ideal_kj": 0.35, "e_max_kj": 0.6}
    ])
    assert req1.input_hash != req2.input_hash


# ---------------------------------------------------------------------------
# RPT-31 … RPT-33: Language and metadata
# ---------------------------------------------------------------------------

def test_rpt31_passport_id_in_csv():
    """RPT-31: Passport ID is written into the CSV output."""
    ctx = build_report_context(
        _typical_specs(),
        passport_id="DWP-TEST-XYZ",
    )
    text = generate_csv_report(ctx).decode("utf-8-sig")
    assert "DWP-TEST-XYZ" in text


def test_rpt32_passport_id_in_xlsx_traceability():
    """RPT-32: Passport ID appears in XLSX Traceability sheet."""
    ctx = build_report_context(
        _typical_specs(),
        passport_id="DWP-PASSPORT-LINK",
    )
    wb = openpyxl.load_workbook(io.BytesIO(generate_xlsx_report(ctx)))
    ws = wb["Traceability"]
    values = [str(c.value or "") for row in ws.iter_rows() for c in row]
    assert "DWP-PASSPORT-LINK" in values


def test_rpt33_app_version_present_in_csv():
    """RPT-33: Application version is included in traceability section of CSV."""
    ctx = _ctx()
    text = generate_csv_report(ctx).decode("utf-8-sig")
    assert ctx.app_version in text


# ---------------------------------------------------------------------------
# RPT-34 … RPT-36: API endpoints
# ---------------------------------------------------------------------------

def test_rpt34_csv_endpoint_requires_auth(client):
    """RPT-34: POST /energy/report/csv returns 401 without auth."""
    r = client.post("/api/v1/energy/report/csv", json=_api_req())
    assert r.status_code == 401


def test_rpt35_xlsx_endpoint_requires_auth(client):
    """RPT-35: POST /energy/report/xlsx returns 401 without auth."""
    r = client.post("/api/v1/energy/report/xlsx", json=_api_req())
    assert r.status_code == 401


def test_rpt36_report_endpoints_accessible_with_auth(client, auth_headers):
    """RPT-36: All three report endpoints return 200 with valid auth and correct content-type."""
    for fmt, ct in [
        ("csv", "text/csv"),
        ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
        ("pdf", "application/pdf"),
    ]:
        r = client.post(f"/api/v1/energy/report/{fmt}", json=_api_req(), headers=auth_headers)
        assert r.status_code == 200, f"Format {fmt}: {r.text}"
        assert ct in r.headers.get("content-type", ""), f"Format {fmt} wrong CT"
        assert len(r.content) > 100, f"Format {fmt} body too short"


# ---------------------------------------------------------------------------
# RPT-37 … RPT-38: Regression guard
# ---------------------------------------------------------------------------

def test_rpt37_existing_energy_dashboard_endpoint_not_broken(client, auth_headers):
    """RPT-37: POST /energy/dashboard still works after P2-F changes."""
    payload = {
        "points": [
            {"point_code": "A", "group_name": "G",
             "e_current_kj": 0.55, "e_min_kj": 0.28, "e_ideal_kj": 0.32, "e_max_kj": 0.42},
        ],
        "production_volume_annual": 100_000,
        "tariff_eur_per_kwh": 0.18,
    }
    r = client.post("/api/v1/energy/dashboard", json=payload, headers=auth_headers)
    assert r.status_code == 200
    body = r.json()
    assert "pareto_detailed" in body
    assert "financial" in body


def test_rpt38_report_csv_content_disposition_header(client, auth_headers):
    """RPT-38: CSV endpoint sets Content-Disposition attachment with .csv filename."""
    r = client.post("/api/v1/energy/report/csv", json=_api_req(), headers=auth_headers)
    assert r.status_code == 200
    cd = r.headers.get("content-disposition", "")
    assert "attachment" in cd
    assert ".csv" in cd
